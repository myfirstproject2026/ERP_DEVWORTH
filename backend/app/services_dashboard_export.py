"""Builds the Dashboard PDF / Excel exports from the same data the Dashboard endpoints return."""
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

EXCEL_TITLE = "DevWorth Dashboard Data"
EXCEL_FILENAME = "DevWorth_Dashboard_Data.xlsx"
PDF_FILENAME = "DevWorth_Dashboard_Data.pdf"

_STATUS_LABELS = {"due_soon": "Due soon", "paid": "Payment received"}


def _status_label(status: str, days_overdue) -> str:
    if status == "overdue":
        return f"{days_overdue or 0} days overdue"
    return _STATUS_LABELS.get(status, status)


def build_excel(company_name: str, summary, pl, production, attendance, followup) -> bytes:
    wb = Workbook()
    wb.properties.title = EXCEL_TITLE
    wb.properties.creator = company_name or "DevWorth"
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1D4ED8")

    def fill_sheet(ws, headers, rows):
        ws.append(headers)
        header_row = ws.max_row
        for cell in ws[header_row]:
            cell.font, cell.fill = head_font, head_fill
        for r in rows or [["No data available"] + [""] * (len(headers) - 1)]:
            ws.append(r)
        for i in range(1, len(headers) + 1):
            col = get_column_letter(i)
            width = max(len(str(c.value or "")) for c in ws[col][header_row - 1:])
            ws.column_dimensions[col].width = min(60, max(14, width + 2))
        return header_row

    ws = wb.active
    ws.title = "Summary"
    ws.append([EXCEL_TITLE])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"{company_name} - generated {datetime.now():%d %b %Y %H:%M}"])
    ws.append([])
    header_row = fill_sheet(ws, ["Metric", "Value", "Detail"], [
        ["Today's sales (INR)", float(summary.todays_sales), f"{summary.sales_change_pct}% vs yesterday"],
        ["Today's purchase (INR)", float(summary.todays_purchase), f"{summary.purchase_change_pct}% vs yesterday"],
        ["Pending payments (INR)", float(summary.pending_payments), f"{summary.overdue_invoice_count} invoices overdue"],
        ["Stock alerts", summary.stock_alert_count, f"{summary.below_reorder_count} items below reorder level"],
        ["AI insight", summary.ai_insight, ""],
    ])
    for row in range(header_row + 1, header_row + 4):
        ws.cell(row=row, column=2).number_format = "#,##0"
        ws.cell(row=row, column=2).alignment = Alignment(horizontal="right")

    fill_sheet(wb.create_sheet("Profit & Loss"), ["Month", "Revenue (lakhs)", "Net profit (lakhs)"],
               [[p.month_label, float(p.revenue_lakhs), float(p.net_profit_lakhs)] for p in pl])
    fill_sheet(wb.create_sheet("Production"), ["Work order", "Product", "Units", "Progress %", "Stage"],
               [[p.wo_number, p.product_name, p.units, p.progress_pct, p.stage] for p in production])
    fill_sheet(wb.create_sheet("Attendance"), ["Present %", "Present", "On leave", "Absent"],
               [[attendance.present_pct, attendance.present_count, attendance.on_leave_count, attendance.absent_count]])
    fill_sheet(wb.create_sheet("Customer Follow-up"),
               ["Customer", "Type", "City", "Amount (INR)", "Status", "Days overdue"],
               [[c.customer_name, c.customer_type, c.city, float(c.amount), c.status, c.days_overdue] for c in followup])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_pdf(company_name: str, user_name: str, summary, pl, production, attendance, followup) -> bytes:
    styles = getSampleStyleSheet()

    def inr(v):
        # built-in PDF fonts have no rupee glyph
        return f"Rs. {float(v):,.0f}"

    def table(headers, rows):
        data = [headers] + (rows or [["No data available"] + [""] * (len(headers) - 1)])
        t = Table(data, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1D4ED8")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#E2E8F0")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ALIGN", (0, 0), (-1, -1), "LEFT"),
        ]))
        return t

    def esc(text):
        return str(text).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    sub = f"Dashboard report - {datetime.now():%A, %d %B %Y}" + (f" - {user_name}" if user_name else "")
    story = [
        Paragraph(esc(company_name or "Dashboard report"), styles["Title"]),
        Paragraph(esc(sub), styles["Normal"]),
        Spacer(1, 8 * mm),
        Paragraph("AI insight", styles["Heading3"]),
        Paragraph(esc(summary.ai_insight), styles["Normal"]),
        Spacer(1, 4 * mm),
        Paragraph("Key metrics", styles["Heading3"]),
        table(["Metric", "Value", "Detail"], [
            ["Today's sales", inr(summary.todays_sales), f"{summary.sales_change_pct}% vs yesterday"],
            ["Today's purchase", inr(summary.todays_purchase), f"{summary.purchase_change_pct}% vs yesterday"],
            ["Pending payments", inr(summary.pending_payments), f"{summary.overdue_invoice_count} invoices overdue"],
            ["Stock alerts", summary.stock_alert_count, f"{summary.below_reorder_count} items below reorder level"],
        ]),
        Spacer(1, 4 * mm),
        Paragraph("Profit &amp; loss - last 6 months (Rs. in lakhs)", styles["Heading3"]),
        table(["Month", "Revenue", "Net profit"], [[p.month_label, p.revenue_lakhs, p.net_profit_lakhs] for p in pl]),
        Spacer(1, 4 * mm),
        Paragraph("Production status", styles["Heading3"]),
        table(["Work order", "Product", "Units", "Progress", "Stage"],
              [[p.wo_number, p.product_name, f"{p.units:,}", f"{p.progress_pct}%", p.stage] for p in production]),
        Spacer(1, 4 * mm),
        Paragraph("Employee attendance (today, all branches)", styles["Heading3"]),
        table(["Present %", "Present", "On leave", "Absent"],
              [[f"{attendance.present_pct}%", attendance.present_count, attendance.on_leave_count, attendance.absent_count]]),
        Spacer(1, 4 * mm),
        Paragraph("Customer follow-up", styles["Heading3"]),
        table(["Customer", "Type / City", "Amount", "Status"],
              [[c.customer_name, f"{c.customer_type} - {c.city}", inr(c.amount), _status_label(c.status, c.days_overdue)]
               for c in followup]),
    ]
    buf = BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, title=EXCEL_TITLE, author=company_name or "DevWorth",
                      leftMargin=18 * mm, rightMargin=18 * mm, topMargin=18 * mm, bottomMargin=18 * mm).build(story)
    return buf.getvalue()
